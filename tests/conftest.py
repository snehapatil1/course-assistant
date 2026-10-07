"""Shared fixtures: synthetic student documents (no course files involved)."""
from __future__ import annotations

from pathlib import Path

import pymupdf
import pytest
from pptx import Presentation

from src import config

# Keep the suite hermetic: even with a real local .env present, tests run with
# all class endpoints switched off unless a test explicitly patches them on.
_ENDPOINT_ATTRS = {
    "chat": ("CHAT_BASE_URL", "CHAT_API_KEY", "CHAT_MODEL"),
    "text_embed": ("TEXT_EMBED_BASE_URL", "TEXT_EMBED_API_KEY", "TEXT_EMBED_MODEL"),
    "visual_embed": ("VISUAL_EMBED_BASE_URL", "VISUAL_EMBED_API_KEY", "VISUAL_EMBED_MODEL"),
    "rerank": ("RERANK_BASE_URL", "RERANK_API_KEY", "RERANK_MODEL"),
    "parse": ("PARSE_BASE_URL", "PARSE_API_KEY", "PARSE_MODEL"),
}


@pytest.fixture(autouse=True)
def _endpoints_off(monkeypatch):
    for attrs in _ENDPOINT_ATTRS.values():
        for a in attrs:
            monkeypatch.setattr(config, a, "")
    yield


@pytest.fixture
def pdf_file(tmp_path: Path) -> Path:
    """A tiny 2-page PDF about RAG (a stand-in for a student's notes)."""
    path = tmp_path / "my_lecture_notes.pdf"
    doc = pymupdf.open()
    for text in (
        "Introduction to retrieval augmented generation.",
        "Retrieval augmented generation combines search with generation for accurate answers.",
    ):
        page = doc.new_page()
        page.insert_text((72, 72), text, fontsize=12)
    doc.save(str(path))
    doc.close()
    return path


@pytest.fixture
def pptx_file(tmp_path: Path) -> Path:
    """A tiny 3-slide deck (a stand-in for a student's slides)."""
    path = tmp_path / "my_deck.pptx"
    prs = Presentation()
    blank = prs.slide_layouts[6]
    for text in ("Slide one: course overview", "Quantization reduces model size",
                 "Few-shot prompting examples"):
        slide = prs.slides.add_slide(blank)
        box = slide.shapes.add_textbox(914400, 914400, 5000000, 1100000)
        box.text_frame.text = text
    prs.save(str(path))
    return path
